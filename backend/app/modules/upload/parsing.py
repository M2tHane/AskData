"""Spreadsheet content is data. No formula evaluation or command interpretation."""

import csv
import math
import re
import unicodedata
from collections.abc import Iterator
from datetime import date, datetime
from io import BytesIO, StringIO
from itertools import islice
from pathlib import PurePath
from typing import Any
from zipfile import BadZipFile, ZipFile

from charset_normalizer import from_bytes
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.core.errors import AppError, ErrorCode
from app.core.importer import ColumnType, ImportColumn
from app.modules.upload.schemas import ParseConfig

MAX_BYTES = 50 * 1024 * 1024


def clean_name(value: str, fallback: str = "column") -> str:
    normalized = unicodedata.normalize("NFKC", value).lower()
    normalized = re.sub(r"[^a-z0-9_]+", "_", normalized).strip("_") or fallback
    if normalized[0].isdigit():
        normalized = "c_" + normalized
    return normalized[:55]


def validate_file(filename: str, data: bytes, max_bytes: int = MAX_BYTES) -> str:
    suffix = PurePath(filename).suffix.lower()
    if (
        not data
        or len(data) > max_bytes
        or PurePath(filename).name != filename
        or "\\" in filename
        or suffix not in {".csv", ".xls", ".xlsx"}
    ):
        raise AppError(ErrorCode.VALIDATION_ERROR, "Invalid spreadsheet file or size")
    if suffix == ".xls" and not data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise AppError(ErrorCode.VALIDATION_ERROR, "Spreadsheet header mismatch")
    if suffix == ".xlsx":
        try:
            with ZipFile(BytesIO(data)) as archive:
                names = archive.namelist()
                if (
                    not {"[Content_Types].xml", "xl/workbook.xml"} <= set(names)
                    or any("vbaproject" in n.lower() for n in names)
                    or sum(info.file_size for info in archive.infolist()) > 250 * 1024 * 1024
                ):
                    raise AppError(ErrorCode.VALIDATION_ERROR, "Unsupported spreadsheet content")
        except BadZipFile:
            raise AppError(ErrorCode.VALIDATION_ERROR, "Spreadsheet header mismatch") from None
    if suffix == ".csv" and (
        b"\x00" in data[:8192] or data.startswith((b"PK\x03\x04", b"\xd0\xcf"))
    ):
        raise AppError(ErrorCode.VALIDATION_ERROR, "CSV header is not text")
    return suffix


def parse_rows(
    filename: str,
    data: bytes,
    config: ParseConfig,
    limit: int | None = None,
) -> tuple[list[str], list[list[Any]], list[str], list[str]]:
    suffix = validate_file(filename, data)
    sheets: list[str] = []
    warnings: list[str] = []
    workbook = None
    source: Iterator[Any]
    try:
        if suffix == ".csv":
            encoding = config.encoding
            if encoding is None:
                detected = from_bytes(data).best()
                encoding = detected.encoding if detected else "utf-8-sig"
            source = csv.reader(StringIO(data.decode(encoding)), delimiter=config.delimiter)
        elif suffix == ".xlsx":
            with ZipFile(BytesIO(data)) as archive:
                if any(
                    b"<mergeCell" in archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("xl/worksheets/") and name.endswith(".xml")
                ):
                    raise AppError(ErrorCode.VALIDATION_ERROR, "请先整理合并单元格或多级表头")
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            sheets = workbook.sheetnames
            sheet = sheets[config.sheet] if isinstance(config.sheet, int) else config.sheet
            source = workbook[sheet].iter_rows(values_only=True)
        else:
            import pandas as pd  # type: ignore[import-untyped]
            from xlrd import open_workbook  # type: ignore

            legacy = open_workbook(file_contents=data, formatting_info=True, on_demand=True)
            try:
                if any(sheet.merged_cells for sheet in legacy.sheets()):
                    raise AppError(ErrorCode.VALIDATION_ERROR, "请先整理合并单元格或多级表头")
            finally:
                legacy.release_resources()
            frames = pd.read_excel(
                BytesIO(data),
                sheet_name=None,
                header=None,
                engine="xlrd",
                dtype=object,
            )
            sheets = list(frames)
            sheet = sheets[config.sheet] if isinstance(config.sheet, int) else config.sheet
            source = iter(frames[sheet].where(frames[sheet].notna(), None).values.tolist())
        if config.header_row is None:
            first = next(source)
            headers = [f"column_{i + 1}" for i in range(len(first))]
            from itertools import chain

            source = chain([first], source)
        else:
            for _ in range(config.header_row - 1):
                next(source)
            headers = [str(v) if v is not None else "" for v in next(source)]
            if any(not h for h in headers):
                warnings.append("表头包含空列；如有多级表头，请先整理")
        rows = [list(row) for row in (islice(source, limit) if limit else source)]
        rows = [row for row in rows if any(v not in (None, "") for v in row)]
        if not headers or any(len(row) != len(headers) for row in rows):
            raise AppError(ErrorCode.VALIDATION_ERROR, "Inconsistent row width")
        return headers, rows, sheets, warnings
    except AppError:
        raise
    except Exception:
        raise AppError(
            ErrorCode.VALIDATION_ERROR, "Cannot parse spreadsheet configuration"
        ) from None
    finally:
        if workbook is not None:
            workbook.close()


def infer_columns(
    headers: list[str], rows: list[list[Any]]
) -> tuple[list[ImportColumn], list[str]]:
    columns, warnings, used = [], [], set()
    for index, original in enumerate(headers):
        name = clean_name(original, f"column_{index + 1}")
        base, number = name, 2
        while name in used:
            name, number = f"{base}_{number}", number + 1
        used.add(name)
        values = [r[index] for r in rows if r[index] not in (None, "")]
        kind: ColumnType = "text"
        texts = [str(v) for v in values]
        if any(re.fullmatch(r"0\d+", t) for t in texts):
            warnings.append(f"{index + 1}: 前导零编码，保留文本")
        elif values and all(isinstance(v, bool) for v in values):
            kind = "boolean"
        elif values and all(isinstance(v, (date, datetime)) for v in values):
            kind = "datetime" if any(isinstance(v, datetime) for v in values) else "date"
        elif values and all(re.fullmatch(r"[+-]?\d+", t) for t in texts):
            kind = "integer"
        elif values:
            try:
                if all(math.isfinite(float(t)) for t in texts):
                    kind = "number"
            except ValueError:
                if any(re.match(r"\d{4}[-/]\d{1,2}[-/]\d", t) for t in texts):
                    warnings.append(f"{index + 1}: 疑似日期文本，请确认类型")
        columns.append(ImportColumn(name=name, data_type=kind, description=original))
    return columns, warnings


def configured_columns(
    config: ParseConfig,
    headers: list[str],
    rows: list[list[Any]],
) -> tuple[list[ImportColumn], list[str]]:
    inferred, warnings = infer_columns(headers, rows)
    if not config.columns:
        return inferred, warnings
    if len(config.columns) != len(headers):
        raise AppError(ErrorCode.VALIDATION_ERROR, "Column count differs")
    result = [
        c.model_copy(
            update={
                "name": clean_name(c.name, f"column_{i + 1}"),
                "description": f"{headers[i]}" + (f" — {c.description}" if c.description else ""),
            }
        )
        for i, c in enumerate(config.columns)
    ]
    if len({c.name for c in result}) != len(result):
        raise AppError(ErrorCode.VALIDATION_ERROR, "Duplicate column names")
    return result, warnings


def convert_rows(rows: list[list[Any]], columns: list[ImportColumn]) -> list[tuple[Any, ...]]:
    def convert(value: Any, column: ImportColumn) -> Any:
        if value in (None, ""):
            return None
        kind = column.data_type
        if kind == "text":
            return str(value)
        if kind == "integer":
            number = int(value)
            if float(value) != number or not -(2**63) <= number < 2**63:
                raise AppError(ErrorCode.VALIDATION_ERROR, "Invalid integer value")
            return number
        if kind == "number":
            number_float = float(value)
            if not math.isfinite(number_float):
                raise AppError(ErrorCode.VALIDATION_ERROR, "Invalid numeric value")
            return number_float
        if kind == "boolean":
            if str(value).lower() not in {"true", "false", "1", "0"}:
                raise AppError(ErrorCode.VALIDATION_ERROR, "Invalid boolean value")
            return str(value).lower() in {"true", "1"}
        if kind == "date":
            return value.date() if isinstance(value, datetime) else date.fromisoformat(str(value))
        return value if isinstance(value, datetime) else datetime.fromisoformat(str(value))

    try:
        return [
            tuple(convert(value, column) for value, column in zip(row, columns, strict=True))
            for row in rows
        ]
    except (ValueError, TypeError, OverflowError):
        raise AppError(ErrorCode.VALIDATION_ERROR, "Column type conversion failed") from None
