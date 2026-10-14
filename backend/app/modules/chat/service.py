import asyncio
import csv
from collections.abc import AsyncIterator, Awaitable, Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import Any, Literal, cast

from fastapi import Request
from fastapi.encoders import jsonable_encoder
from openpyxl import Workbook  # type: ignore[import-untyped]

from app.core.audit import (
    AuditService,
    DatabaseAuditStore,
    DatabaseQueryLogStore,
    DetachedQueryLogStore,
    QueryLogWriter,
)
from app.core.cache import TextCache
from app.core.datasource import Datasource
from app.core.db import AsyncSession, Database
from app.core.deps import CurrentUser
from app.core.errors import MESSAGES, AppError, ErrorCode, SqlBlocked
from app.core.llm import CallContext, CallRecord, LLMClient, RerankClient, UsageRecorder
from app.core.masking import mask_rows, safe_export_cell
from app.core.permission import VisibleScope
from app.core.response import Page, PageParams
from app.core.sql.executor import QueryLimits, SqlExecutor
from app.core.sql.guard import GuardedSql, SqlGuard
from app.core.stream import EventType, StreamEvent
from app.modules.chat.pipeline import (
    PREVIEW_ROWS,
    ContextBuilder,
    ExecuteStage,
    GuardStage,
    PipelineDependencies,
    QueryContext,
    QueryPipeline,
    RecordStage,
    RetrievalIndex,
    diagnose,
    encode_event,
)
from app.modules.chat.pipeline import ScopeOnlyIndex as ScopeOnlyIndex
from app.modules.chat.pipeline import plain_rows as plain_rows
from app.modules.chat.repository import ChatRepository
from app.modules.chat.retrieval import ServiceRetrievalIndex
from app.modules.chat.schemas import (
    ChatLimits,
    ContextItemView,
    ContextPreview,
    ContextPreviewRequest,
    ConversationRename,
    ConversationView,
    ExportFile,
    MessageView,
    QueryReply,
    RerunRequest,
    ResultSnapshot,
    SuggestedQuestions,
)
from app.modules.chat.schemas import (
    ConversationCreate as ConversationCreate,
)
from app.modules.chat.schemas import (
    QuestionRequest as QuestionRequest,
)
from app.modules.chat.suggestions import SuggestionService
from app.modules.querylog.service import build_service as build_querylog

# Exports of results larger than the kept preview re-run the SQL under fresh grants, with caps
# larger than the interactive ones.
EXPORT_MAX_ROWS = 100_000
EXPORT_MAX_BYTES = 50 * 1024 * 1024
EXPORT_TIMEOUT_SECONDS = 120.0
# ChatService has a list() method: annotations inside it use these names instead.
type Columns = list[str]
type Rows = list[dict[str, Any]]
type Kept = tuple[dict[str, list[str]], list[str]]


@dataclass
class Review:
    """An earlier answer checked against today's grants."""

    revoked: bool = False
    # "table.column" the user can no longer read; empty when they cannot be named.
    lost: list[str] = field(default_factory=list)
    # A result column is masked now that was not when answered.
    masks_changed: bool = False
    guarded: GuardedSql | None = None


def lost_fields(sources: dict[str, list[str]], scope: VisibleScope) -> list[str]:
    lost = []
    for table, columns in sorted(sources.items()):
        visible = scope.tables.get(table)
        name = table.rsplit(".", 1)[-1]
        lost += [f"{name}.{c}" for c in sorted(columns) if not visible or c not in visible.columns]
    return lost


def snapshot_of(ctx: QueryContext, *, reread: bool = False) -> ResultSnapshot | None:
    result, guarded = ctx.result, ctx.guarded_sql
    if ctx.status != "success" or result is None or guarded is None:
        return None
    return ResultSnapshot(
        columns=result.columns,
        # Already masked by the executor; JSON-safe like the reply itself.
        rows=jsonable_encoder(plain_rows(result.rows[:PREVIEW_ROWS])),
        row_count=result.row_count,
        truncated=result.truncated,
        preview_truncated=len(result.rows) > PREVIEW_ROWS,
        sources={table: sorted(columns) for table, columns in guarded.columns.items()},
        masked=sorted(guarded.output_masks),
        queried_at=datetime.now(UTC),
        reread=reread,
    )


# New conversations start with this title until the first answer names them.
DEFAULT_TITLE = ConversationCreate.model_fields["title"].default


# The user bubble of an edited-SQL rerun; the real question travels in rerun_of.
RERUN_QUESTION = "手动编辑 SQL 重跑"


class ChatService:
    """Request-scoped service. Repository, audit and querylog share one DB session.

    Writes commit with their audit records before SSE completion. Do not put a
    rollback-on-query-error outer transaction around this service: failed runs are data.
    """

    def __init__(
        self,
        repository: ChatRepository,
        pipeline: QueryPipeline,
        audit: AuditService,
        cache: TextCache | None = None,
        *,
        export_max_rows: int = EXPORT_MAX_ROWS,
        export_max_bytes: int = EXPORT_MAX_BYTES,
    ) -> None:
        self.repository, self.pipeline, self.audit = repository, pipeline, audit
        self.export_max_rows, self.export_max_bytes = export_max_rows, export_max_bytes
        self.suggestions = SuggestionService(pipeline.deps, cache)
        self.queries = (
            build_querylog(repository.session) if isinstance(repository, ChatRepository) else None
        )

    async def create(self, request: ConversationCreate, *, actor: CurrentUser) -> ConversationView:
        await self.pipeline.deps.permissions.visible_scope(actor, request.datasource_id)
        await self.pipeline.deps.datasource(request.datasource_id, actor)
        try:
            conversation = await self.repository.create(actor.id, actor.team_id, request)
            await self.audit.record(actor, "chat.create", ("conversation", str(conversation.id)))
            await self.repository.commit()
        except Exception:
            await self.repository.rollback()
            raise AppError(ErrorCode.INTERNAL_ERROR, http_status=500) from None
        return conversation

    async def list(
        self, page: PageParams, *, actor: CurrentUser, query: str = ""
    ) -> Page[ConversationView]:
        items, total = await self.repository.list_conversations(actor.id, page, query.strip())
        return Page(items=items, total=total, page=page.page, page_size=page.page_size)

    async def get(self, identifier: int, *, actor: CurrentUser) -> ConversationView:
        return await self.repository.get(identifier, actor.id)

    async def messages(
        self, identifier: int, page: PageParams, *, actor: CurrentUser
    ) -> Page[MessageView]:
        conversation = await self.repository.get(identifier, actor.id)
        items, total = await self.repository.messages(identifier, page)
        answers = [m.id for m in items if m.role == "assistant" and m.normalized_sql]
        if answers:
            kept = await self.repository.result_sources(answers)
            grants = await self._grants(conversation, actor)
            items = [
                self._withhold(m, self._review(m.normalized_sql, kept.get(m.id), grants))
                if m.id in answers and m.normalized_sql
                else m
                for m in items
            ]
        return Page(items=items, total=total, page=page.page, page_size=page.page_size)

    async def _grants(
        self, conversation: ConversationView, actor: CurrentUser
    ) -> tuple[VisibleScope, str] | None:
        """Today's scope and dialect; None when the datasource is no longer usable."""
        deps = self.pipeline.deps
        try:
            datasource = await deps.datasource(conversation.datasource_id, actor)
            scope = await deps.permissions.visible_scope(actor, conversation.datasource_id)
        except AppError:
            return None
        return scope, datasource.dialect

    def _review(
        self,
        sql: str,
        kept: Kept | None,
        grants: tuple[VisibleScope, str] | None,
    ) -> Review:
        """May this earlier answer still be shown? kept = (sources, masked) of its rows."""
        if grants is None:
            return Review(revoked=True)
        scope, dialect = grants
        try:
            guarded = self.pipeline.deps.guard.check(sql, dialect, scope)
        except AppError:
            return Review(revoked=True, lost=lost_fields(kept[0], scope) if kept else [])
        # Answers from before results were kept: what was masked then is unknown.
        changed = kept is not None and bool(set(guarded.output_masks) - set(kept[1]))
        return Review(masks_changed=changed, guarded=guarded)

    @staticmethod
    def _withhold(message: MessageView, review: Review) -> MessageView:
        if review.revoked:
            # Nothing of the answer: its text, reasoning and chart came from the lost fields.
            return message.model_copy(
                update={
                    "content": "",
                    "meta": None,
                    "chart": None,
                    "normalized_sql": None,
                    "revoked": True,
                    "revoked_fields": review.lost,
                }
            )
        if review.masks_changed:
            # The text and reasoning may quote values that are masked now.
            meta = message.meta.model_copy(update={"reasoning": ""}) if message.meta else None
            return message.model_copy(update={"content": "", "meta": meta, "masks_changed": True})
        return message

    async def rename(
        self, identifier: int, request: ConversationRename, *, actor: CurrentUser
    ) -> ConversationView:
        try:
            conversation = await self.repository.rename(identifier, actor.id, request.title)
            await self.audit.record(actor, "chat.rename", ("conversation", str(identifier)))
            await self.repository.commit()
        except AppError:
            await self.repository.rollback()
            raise
        return conversation

    async def delete(self, identifier: int, *, actor: CurrentUser) -> None:
        await self.repository.get(identifier, actor.id)
        await self.repository.delete(identifier, actor.id)
        await self.audit.record(actor, "chat.delete", ("conversation", str(identifier)))
        await self.repository.commit()

    async def prepare(
        self,
        identifier: int,
        request: QuestionRequest,
        *,
        actor: CurrentUser,
        sql: str | None = None,
    ) -> QueryContext:
        conversation = await self.repository.get(identifier, actor.id)
        question = request.question
        query_id = None
        used = False
        ref = None
        # The user's bubble shows what they typed or the option they picked, never a mix.
        shown = question
        if request.clarification_message_id is not None:
            if not request.clarification_choice:
                raise AppError(ErrorCode.VALIDATION_ERROR)
            query_id = await self.repository.clarification_query_id(
                identifier, request.clarification_message_id
            )
            # The pipeline continues the original question; the choice travels separately.
            question, ref = await self.repository.consume_clarification(
                identifier, request.clarification_message_id, request.clarification_choice
            )
            shown = request.clarification_choice
            used = True
        elif request.clarification_choice:
            raise AppError(ErrorCode.VALIDATION_ERROR)
        message = await self.repository.add_message(identifier, actor.id, "user", shown)
        await self.audit.record(
            actor, "chat.rerun" if sql else "chat.ask", ("message", str(message.id))
        )
        await self.repository.commit()
        return QueryContext(
            user=actor,
            datasource_id=conversation.datasource_id,
            conversation_id=identifier,
            question=question,
            timezone=request.timezone,
            manual_sql=sql,
            clarification_used=used,
            clarification_choice=request.clarification_choice,
            clarification_ref=ref,
            pinned_metric_ids=frozenset(request.referenced_metric_ids),
            requested_model_id=getattr(request, "model_id", None),
            message_id=message.id,
            query_id=query_id or f"chat:{identifier}:{message.id}",
        )

    async def finish(self, ctx: QueryContext) -> QueryReply:
        query_created_at = None
        if self.queries is not None and ctx.query_run_id is not None:
            query_created_at = (await self.queries.get(ctx.query_run_id, ctx.user)).created_at
        content = ctx.clarification.question if ctx.clarification else ctx.explanation
        message = await self.repository.add_message(
            ctx.conversation_id,
            ctx.user.id,
            "assistant",
            content,
            query_run_id=ctx.query_run_id,
            sql=ctx.guarded_sql.normalized_sql
            if ctx.guarded_sql and ctx.status == "success"
            else None,
            clarification=ctx.clarification,
            parent_message_id=ctx.message_id,
            query_run_created_at=query_created_at,
            chart=ctx.chart if ctx.status == "success" else None,
            meta=ctx.answer_meta() if ctx.status in {"success", "clarify"} else None,
        )
        snapshot = snapshot_of(ctx)
        if snapshot is not None:
            await self.repository.save_result(ctx.conversation_id, message.id, snapshot)
        await self.audit.record(
            ctx.user, "chat.answer", ("message", str(message.id)), {"status": ctx.status}
        )
        title = await self._name_conversation(ctx)
        await self.repository.commit()
        reply = ctx.reply()
        reply.message_id = message.id
        reply.title = title
        return reply

    async def _name_conversation(self, ctx: QueryContext) -> str | None:
        """The first question names a conversation that still has the default title.

        Failed and blocked questions name it too, so the list never fills up with "新会话".
        """
        if ctx.manual_sql:
            return None
        title = ctx.title or " ".join(ctx.question.split())[:20]
        if not title or not await self.repository.rename_default(
            ctx.conversation_id, DEFAULT_TITLE, title
        ):
            return None
        return title

    async def ask(
        self, identifier: int, request: QuestionRequest, *, actor: CurrentUser
    ) -> QueryReply:
        ctx = await self.prepare(identifier, request, actor=actor)
        await self.pipeline.run(ctx)
        await self.repository.commit()
        return await self.finish(ctx)

    async def run_question(
        self,
        ctx: QueryContext,
        emit: Callable[[EventType, dict[str, Any]], Awaitable[None]] | None = None,
    ) -> QueryReply:
        """Answer a prepared question without streaming (MCP): ask() with progress events."""
        ctx.emit = emit
        try:
            await self.pipeline.run(ctx)
        except AppError as error:
            if error.code == ErrorCode.QUERY_CANCELLED:
                await self._commit_cancelled()
            raise
        finally:
            ctx.emit = None
        await self.repository.commit()
        return await self.finish(ctx)

    async def execute_sql(
        self, datasource_id: int, sql: str, *, actor: CurrentUser, origin: dict[str, Any]
    ) -> QueryContext:
        """Run SQL the caller wrote (MCP run_sql) under the caller's grants and masks.

        The same guard and executor as a generated query; logged, but kept in no conversation.
        """
        deps = self.pipeline.deps
        ctx = QueryContext(
            user=actor,
            datasource_id=datasource_id,
            conversation_id=0,
            question="执行 SQL",
            manual_sql=sql,
            generated_sql=sql,
            origin=origin,
        )
        try:
            ctx.scope = await deps.permissions.visible_scope(actor, datasource_id)
            ctx.datasource = await deps.datasource(datasource_id, actor)
            for stage in (GuardStage(deps), ExecuteStage(deps)):
                ctx.stage = stage.name
                await stage.run(ctx)
        except AppError as error:
            ctx.diagnostic = diagnose(ctx, error)
            ctx.error_code = error.code
            ctx.status = (
                "blocked"
                if isinstance(error, SqlBlocked)
                or error.code.value.startswith("SQL_")
                and error.code != ErrorCode.SQL_EXECUTION_ERROR
                else "failed"
            )
        await RecordStage(deps).run(ctx)
        await self.repository.commit()
        return ctx

    async def rerun(
        self, identifier: int, request: RerunRequest, *, actor: CurrentUser
    ) -> QueryReply:
        ctx = await self.prepare(
            identifier,
            QuestionRequest(question=RERUN_QUESTION, timezone=request.timezone),
            actor=actor,
            sql=request.sql,
        )
        if request.source_message_id is not None:
            # No model call writes a chart for edited SQL: keep the one the user already saw.
            # Presentation validates it against the new columns and falls back by shape.
            source = await self.repository.assistant_message(identifier, request.source_message_id)
            if source.chart is not None:
                ctx.chart_hint = source.chart.model_dump()
            # Older answers carry no question in their meta; their user message holds it.
            question = (source.meta.question if source.meta else "") or (
                await self.repository.asked_before(identifier, source.id)
            )
            ctx.rerun_of = {
                "message_id": source.id,
                "query_run_id": source.query_run_id,
                "question": "" if question == RERUN_QUESTION else question,
            }
        await self.pipeline.run(ctx)
        await self.repository.commit()
        return await self.finish(ctx)

    async def replay(self, identifier: int, message_id: int, *, actor: CurrentUser) -> QueryReply:
        """An earlier answer's rows: the kept ones once today's grants allow them.

        Answers from before results were kept are queried once more, and that result is kept.
        """
        conversation = await self.repository.get(identifier, actor.id)
        message = await self.repository.assistant_message(identifier, message_id)
        if not message.normalized_sql:
            raise AppError(ErrorCode.NOT_FOUND, http_status=404)
        snapshot = await self.repository.result(message_id)
        if snapshot is not None:
            review = self._review(
                message.normalized_sql,
                (snapshot.sources, snapshot.masked),
                await self._grants(conversation, actor),
            )
            if review.revoked or review.guarded is None:
                return QueryReply(
                    status="blocked",
                    message_id=message_id,
                    code=ErrorCode.SQL_OUT_OF_SCOPE,
                    revoked_fields=review.lost,
                )
            changed = review.masks_changed
            return QueryReply(
                status="success",
                message_id=message_id,
                sql=message.normalized_sql,
                columns=snapshot.columns,
                rows=mask_rows(snapshot.rows, review.guarded.output_masks),
                truncated=snapshot.truncated,
                row_count=snapshot.row_count,
                preview_truncated=snapshot.preview_truncated,
                chart=message.chart,
                explanation="" if changed else message.content,
                meta=message.meta.model_copy(update={"reasoning": ""}) if message.meta else None,
                queried_at=snapshot.queried_at,
                reread=snapshot.reread,
                masks_changed=changed,
            )
        ctx = QueryContext(
            user=actor,
            datasource_id=conversation.datasource_id,
            conversation_id=identifier,
            question="查看历史结果",
            message_id=message_id,
            manual_sql=message.normalized_sql,
            generated_sql=message.normalized_sql,
            limit_overrides={"purpose": "replay"},
        )
        try:
            ctx.scope = await self.pipeline.deps.permissions.visible_scope(
                actor, conversation.datasource_id
            )
            ctx.datasource = await self.pipeline.deps.datasource(conversation.datasource_id, actor)
            await GuardStage(self.pipeline.deps).run(ctx)
            await ExecuteStage(self.pipeline.deps).run(ctx)
        except AppError as error:
            if error.code == ErrorCode.RATE_LIMITED:
                # Busy, not an outcome of the answer: the client waits and reads again.
                raise
            blocked = error.code == ErrorCode.FORBIDDEN or (
                error.code.value.startswith("SQL_") and error.code != ErrorCode.SQL_EXECUTION_ERROR
            )
            # Access revoked: nothing of the answer. Otherwise (timeout, datasource down) the
            # answer itself still stands and only its rows could not be read.
            return QueryReply(
                status="blocked" if blocked else "failed",
                message_id=message_id,
                code=error.code,
                explanation="" if blocked else message.content,
            )
        ctx.status = "success"
        snapshot = snapshot_of(ctx, reread=True)
        if snapshot is not None:
            await self.repository.save_result(identifier, message_id, snapshot)
            await self.repository.commit()
        ctx.chart = message.chart
        ctx.explanation = message.content
        ctx.meta = message.meta.model_copy(update={"reasoning": ""}) if message.meta else None
        reply = ctx.reply()
        reply.queried_at = snapshot.queried_at if snapshot else None
        reply.reread = True
        return reply

    async def _commit_cancelled(self) -> None:
        # A detached writer already committed the run; the request session may be closing.
        if self.pipeline.deps.cancelled_querylog is None:
            await self.repository.commit()

    def stream(self, ctx: QueryContext) -> AsyncIterator[StreamEvent]:
        return self._stream(ctx)

    async def _stream(self, ctx: QueryContext) -> AsyncIterator[StreamEvent]:
        queue: asyncio.Queue[StreamEvent | None] = asyncio.Queue(maxsize=8)
        identifier = 0

        async def emit(kind: EventType, data: dict[str, Any]) -> None:
            nonlocal identifier
            if kind == EventType.DONE:
                return
            await queue.put(encode_event(identifier, kind, jsonable_encoder(data)))
            identifier += 1

        async def produce() -> None:
            nonlocal identifier
            ctx.emit = emit
            try:
                await self.pipeline.run(ctx)
                await self.repository.commit()
                reply = await self.finish(ctx)
                if reply.title:
                    await queue.put(
                        encode_event(identifier, EventType.TITLE, {"title": reply.title})
                    )
                    identifier += 1
                await queue.put(
                    encode_event(
                        identifier,
                        EventType.DONE,
                        {
                            "status": reply.status,
                            "query_run_id": reply.query_run_id,
                            "message_id": reply.message_id,
                        },
                    )
                )
            except AppError as error:
                if error.code == ErrorCode.QUERY_CANCELLED:
                    # Disconnect: the run is already logged; no queue I/O.
                    await self._commit_cancelled()
                    return
                await self.repository.rollback()
                await queue.put(
                    encode_event(
                        identifier,
                        EventType.ERROR,
                        {
                            "code": error.code.value,
                            "message": MESSAGES[error.code],
                        },
                    )
                )
                identifier += 1
                await queue.put(encode_event(identifier, EventType.DONE, {"status": "failed"}))
            except asyncio.CancelledError:
                await self._commit_cancelled()
                return
            except Exception:
                await self.repository.rollback()
                await queue.put(
                    encode_event(
                        identifier,
                        EventType.ERROR,
                        {
                            "code": ErrorCode.INTERNAL_ERROR.value,
                            "message": MESSAGES[ErrorCode.INTERNAL_ERROR],
                        },
                    )
                )
                identifier += 1
                await queue.put(encode_event(identifier, EventType.DONE, {"status": "failed"}))
            finally:
                task = asyncio.current_task()
                if task is not None and not task.cancelling():
                    await queue.put(None)

        producer = asyncio.create_task(produce())
        try:
            while True:
                event = await queue.get()
                if event is None:
                    break
                yield event
        finally:
            # Disable emission before cancellation so a full queue cannot deadlock cleanup.
            ctx.emit = None
            if not producer.done():
                producer.cancel()
            await asyncio.gather(producer, return_exceptions=True)

    async def export(
        self,
        identifier: int,
        message_id: int,
        format: Literal["csv", "xlsx"],
        *,
        actor: CurrentUser,
    ) -> ExportFile:
        conversation = await self.repository.get(identifier, actor.id)
        sql_text = await self.repository.execution_sql(identifier, message_id)
        snapshot = await self.repository.result(message_id)
        if snapshot is not None and snapshot.complete:
            # Every row is kept: the file matches what the answer showed, no new query.
            review = self._review(
                sql_text,
                (snapshot.sources, snapshot.masked),
                await self._grants(conversation, actor),
            )
            await self.audit.record(
                actor,
                "chat.export",
                ("message", str(message_id)),
                {
                    "format": format,
                    "status": "blocked" if review.guarded is None else "success",
                    "row_count": 0 if review.guarded is None else snapshot.row_count,
                    "source": "kept",
                },
            )
            await self.repository.commit()
            if review.guarded is None:
                raise AppError(ErrorCode.SQL_OUT_OF_SCOPE)
            columns = snapshot.columns
            rows = mask_rows(snapshot.rows[: self.export_max_rows], review.guarded.output_masks)
            truncated = len(snapshot.rows) > self.export_max_rows
        else:
            columns, rows, truncated = await self._export_query(
                conversation, message_id, sql_text, format, actor
            )
        return self._export_file(message_id, format, columns, rows, truncated)

    async def _export_query(
        self,
        conversation: ConversationView,
        message_id: int,
        sql_text: str,
        format: Literal["csv", "xlsx"],
        actor: CurrentUser,
    ) -> tuple[Columns, Rows, bool]:
        """A new query for the whole result, under fresh grants, masks and export caps."""
        identifier = conversation.id
        ctx = QueryContext(
            user=actor,
            datasource_id=conversation.datasource_id,
            conversation_id=identifier,
            question="导出查询结果",
            message_id=message_id,
            manual_sql=sql_text,
            generated_sql=sql_text,
            # Exports re-run the guarded SQL with the larger export caps, never the preview cap.
            limit_overrides={
                "max_rows": self.export_max_rows,
                "timeout_seconds": EXPORT_TIMEOUT_SECONDS,
            },
        )
        try:
            ctx.scope = await self.pipeline.deps.permissions.visible_scope(
                actor, conversation.datasource_id
            )
            ctx.datasource = await self.pipeline.deps.datasource(conversation.datasource_id, actor)
            await GuardStage(self.pipeline.deps).run(ctx)
            await ExecuteStage(self.pipeline.deps).run(ctx)
        except AppError as error:
            ctx.error_code = error.code
            ctx.status = (
                "blocked"
                if error.code.value.startswith("SQL_")
                and (error.code != ErrorCode.SQL_EXECUTION_ERROR)
                else "failed"
            )
            raise AppError(error.code, http_status=error.http_status) from None
        finally:
            await self.pipeline.stages[-1].run(ctx)
            await self.audit.record(
                actor,
                "chat.export",
                ("message", str(message_id)),
                {
                    "format": format,
                    "status": ctx.status,
                    "row_count": ctx.result.row_count if ctx.result else 0,
                    "source": "query",
                },
            )
            await self.repository.commit()
        result, guarded = ctx.result, ctx.guarded_sql
        if result is None or guarded is None:
            raise AppError(ErrorCode.INTERNAL_ERROR)
        return result.columns, mask_rows(result.rows, guarded.output_masks), result.truncated

    def _export_file(
        self,
        message_id: int,
        format: Literal["csv", "xlsx"],
        columns: Columns,
        rows: Rows,
        truncated: bool,
    ) -> ExportFile:
        def export_cell(value: Any) -> Any:
            if isinstance(value, datetime) and value.tzinfo is not None:
                value = value.isoformat()
            elif isinstance(value, (dict, list, tuple, bytes)):
                value = str(value)
            return safe_export_cell(value)

        cells = [[export_cell(column) for column in columns]]
        # Size cap: estimate the encoded size row by row and stop before the file exceeds it.
        size = 0
        for record in rows:
            line = [export_cell(record.get(column)) for column in columns]
            size += sum(len(str(value).encode()) + 1 for value in line)
            if size > self.export_max_bytes:
                truncated = True
                break
            cells.append(line)
        if format == "csv":
            output = StringIO(newline="")
            csv.writer(output).writerows(cells)
            content = output.getvalue().encode("utf-8-sig")
            content_type = "text/csv; charset=utf-8"
        else:
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("查询结果")
            for row in cells:
                sheet.append(row)
            buffer = BytesIO()
            workbook.save(buffer)
            content = buffer.getvalue()
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return ExportFile(
            filename=f"query_{message_id}.{format}",
            content_type=content_type,
            content=content,
            row_count=len(cells) - 1,
            truncated=truncated,
        )

    def limits(self) -> ChatLimits:
        return ChatLimits(
            preview_rows=PREVIEW_ROWS,
            export_max_rows=self.export_max_rows,
            export_max_mb=self.export_max_bytes // (1024 * 1024),
        )

    async def context_preview(
        self, request: ContextPreviewRequest, *, actor: CurrentUser
    ) -> ContextPreview:
        deps = self.pipeline.deps
        scope = await deps.permissions.visible_scope(actor, request.datasource_id)
        await deps.datasource(request.datasource_id, actor)
        if not request.question and len(scope.tables) > deps.builder.threshold:
            # Large datasets retrieve by question; there is nothing meaningful to preview.
            raise AppError(ErrorCode.VALIDATION_ERROR)
        bundle = await deps.builder.build(
            actor, scope, request.question, None, frozenset(request.referenced_metric_ids)
        )
        return ContextPreview(
            mode=bundle.mode,
            tables=bundle.tables,
            items=[
                ContextItemView(
                    kind=item.kind,
                    name=item.name,
                    scope=item.scope.value,
                    version=item.version,
                    definition=item.definition,
                    pinned=item.key in bundle.pinned,
                )
                for item in bundle.items
            ],
            required_conditions=bundle.required_conditions,
            ambiguities=bundle.ambiguities,
        )

    async def suggested_questions(
        self, datasource_id: int, *, actor: CurrentUser
    ) -> SuggestedQuestions:
        return await self.suggestions.suggest(actor, datasource_id)


class ChatRetentionStore:
    """Compose with core.retention; delete only this module's expired messages."""

    def __init__(self, repository: ChatRepository) -> None:
        self.repository = repository

    async def purge(self, before: datetime) -> dict[str, int]:
        return {"message": await self.purge_before(before)}

    async def purge_before(self, now: datetime) -> int:
        return await self.repository.purge_before(now)


def build_retention_store(session: AsyncSession) -> ChatRetentionStore:
    return ChatRetentionStore(ChatRepository(session))


class _MissingUsageRecorder:
    """Fail closed until the quota/usage service is wired through its public protocol."""

    async def authorize_query(self, context: CallContext) -> None:
        raise AppError(ErrorCode.LLM_UNAVAILABLE, http_status=503)

    async def authorize(self, context: CallContext) -> None:
        raise AppError(ErrorCode.LLM_UNAVAILABLE, http_status=503)

    async def record(self, record: CallRecord) -> None:
        return None


def build_service(session: AsyncSession, request: Request) -> ChatService:
    from app.modules.auth.service import secret_box, settings_for
    from app.modules.datasource.service import build_service as build_datasource
    from app.modules.feedback.service import build_service as build_feedback
    from app.modules.metric.service import QueryDatasources
    from app.modules.metric.service import build_service as build_metric
    from app.modules.modelcfg.service import build_service as build_models
    from app.modules.permission.service import build_service as build_permissions
    from app.modules.promptcfg.service import build_service as build_prompts
    from app.modules.schema.service import build_service as build_schema
    from app.modules.semantic.service import SchemaTargets
    from app.modules.semantic.service import build_service as build_semantic

    settings = settings_for(request)
    box = secret_box(settings)
    permissions = build_permissions(
        session, box, getattr(request.app.state, "scope_cache", None)
    ).permissions
    datasource_service = build_datasource(session, box)

    async def resolve(identifier: int, actor: CurrentUser) -> Datasource:
        return await datasource_service.for_query(identifier, actor, permissions)

    audit = AuditService(DatabaseAuditStore(session), settings.retention_days)
    repository = ChatRepository(session)
    recorder = cast(
        UsageRecorder, getattr(request.app.state, "usage_recorder", _MissingUsageRecorder())
    )
    models = build_models(session, box)
    llm = cast(LLMClient, getattr(request.app.state, "llm_client", None))
    if llm is None:
        llm = LLMClient(
            models.registry, recorder, getattr(request.app.state, "model_transport", None)
        )
    index = cast(RetrievalIndex | None, getattr(request.app.state, "retrieval_index", None))
    executor = cast(SqlExecutor, getattr(request.app.state, "sql_executor", None))
    if executor is None:
        # Core alone constructs connectors and distributed rate limiting.
        cache = getattr(request.app.state, "scope_cache", None)
        executor = SqlExecutor.default(getattr(cache, "redis", None))
    if index is None:
        schema = build_schema(session, box)
        resolver = QueryDatasources(datasource_service, permissions)
        semantic = build_semantic(session, permissions, SchemaTargets(schema))
        metric = build_metric(session, permissions, resolver, executor)
        feedback = build_feedback(session, build_querylog(session), permissions, resolver)
        index = ServiceRetrievalIndex(
            schema,
            semantic,
            metric,
            feedback,
            getattr(request.state, "embedding_service", None),
            RerankClient(
                models.registry, recorder, getattr(request.app.state, "model_transport", None)
            ),
        )
    database = cast(Database | None, getattr(request.app.state, "database", None))
    cancelled_log = (
        QueryLogWriter(DetachedQueryLogStore(database.sessions), settings.retention_days)
        if database is not None
        else None
    )
    dependencies = PipelineDependencies(
        permissions,
        resolve,
        repository.history,
        ContextBuilder(index, threshold=0 if isinstance(index, ServiceRetrievalIndex) else 30),
        llm,
        SqlGuard(),
        executor,
        QueryLogWriter(DatabaseQueryLogStore(session), settings.retention_days),
        prompt_source=build_prompts(session),
        cancelled_querylog=cancelled_log,
        limits=QueryLimits(
            user_id=1,
            timeout_seconds=settings.query_timeout_seconds,
            max_rows=settings.max_query_rows,
            max_cost=settings.max_query_cost,
        ),
    )
    return ChatService(
        repository,
        QueryPipeline(dependencies),
        audit,
        export_max_rows=settings.export_max_rows,
        export_max_bytes=settings.export_max_mb * 1024 * 1024,
    )


async def get_service(request: Request) -> AsyncIterator[ChatService]:
    injected = getattr(request.state, "chat_service", None)
    if isinstance(injected, ChatService):
        yield injected
        return
    database = cast(Database | None, getattr(request.app.state, "database", None))
    if database is None:
        raise AppError(ErrorCode.DATASOURCE_UNAVAILABLE, http_status=503)
    # Request scope keeps this session open through the complete SSE body.
    async with database.session() as session:
        yield build_service(session, request)
