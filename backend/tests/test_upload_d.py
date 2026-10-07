import base64
import importlib.util
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import Workbook
from pydantic import SecretStr
from sqlalchemy import Column, Integer, MetaData, Table

from app.core.audit import AuditLog, AuditService, DatabaseAuditStore
from app.core.datasource import Datasource
from app.core.db import Database
from app.core.deps import CurrentUser, Role
from app.core.errors import AppError, ErrorCode
from app.core.importer import ImportColumn, ImportPlan, PostgresImportGateway
from app.core.permission import PermissionService, TableGrant
from app.core.sql.dialect import Dialect
from app.core.tasks import TaskManager, TaskStatus
from app.modules.datasource import models as datasource_models  # noqa: F401
from app.modules.org import models as org_models  # noqa: F401
from app.modules.upload.parsing import configured_columns, convert_rows, parse_rows, validate_file
from app.modules.upload.repository import UploadRepository
from app.modules.upload.schemas import ConfirmRequest, ParseConfig, PreviewRequest
from app.modules.upload.service import UploadService
from tests.fakes import FakePermissions, MemoryQueue, MemoryTaskStore

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
async def d_database(tmp_path):
    database = Database(f"sqlite+aiosqlite:///{tmp_path / 'd.db'}")
    spec = importlib.util.spec_from_file_location(
        "d_migration",
        ROOT / "backend/migrations/versions/0005_upload_eval.py",
    )
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)

    def initialize(connection):
        # Only external FK targets are placeholders; D schema uses the actual migration.
        metadata = MetaData()
        for name in ("datasource", "team"):
            Table(name, metadata, Column("id", Integer, primary_key=True))
        metadata.create_all(connection)
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        AuditLog.__table__.create(connection)

    async with database.engine.begin() as connection:
        await connection.run_sync(initialize)
    yield database
    await database.close()


class Storage:
    def __init__(self):
        self.values = {}

    async def put(self, key, data, content_type):
        import hashlib

        self.values[key] = data
        return SimpleNamespace(sha256=hashlib.sha256(data).hexdigest(), size=len(data))

    async def get(self, key):
        return self.values[key]

    async def delete(self, key):
        self.values.pop(key, None)


class Registration:
    def __init__(self):
        self.published = []
        self.fail = False

    async def reserve(self, name, team_id, actor):
        return 7

    async def owner_team(self, datasource_id):
        return 1

    async def publish(self, job_id, datasource, table, owner_team_id, actor):
        if self.fail:
            raise AppError(ErrorCode.DATASOURCE_UNAVAILABLE)
        self.published.append((job_id, datasource, table, owner_team_id))


class Importer:
    def __init__(self):
        self.tables = {}
        self.writes = 0
        self.fail = False

    def query_datasource(self, datasource_id):
        return Datasource(
            id=datasource_id,
            dialect=Dialect.POSTGRES,
            host="pg-import",
            port=5432,
            database="imports",
            username=f"ds_{datasource_id}_reader",
            password=SecretStr("query-secret"),
            schema_name=f"ds_{datasource_id}",
        )

    @asynccontextmanager
    async def transaction(self, plan):
        previous = self.tables.get(plan.table_name)
        structure = [(c.name, c.data_type) for c in plan.columns]
        if plan.mode == "append" and (previous is None or previous[0] != structure):
            raise AppError(ErrorCode.CONFLICT)
        importer = self

        class Writer:
            async def write(self, rows):
                importer.writes += 1
                if importer.fail:
                    raise AppError(ErrorCode.DATASOURCE_UNAVAILABLE)
                importer.tables[plan.table_name] = (
                    structure,
                    (previous[1] if plan.mode == "append" else []) + list(rows),
                )
                return len(rows)

        try:
            yield Writer()
        except Exception:
            if previous is None:
                self.tables.pop(plan.table_name, None)
            else:
                self.tables[plan.table_name] = previous
            raise


@pytest.fixture
async def upload_setup(d_database):
    storage, registration, importer = Storage(), Registration(), Importer()
    store, queue = MemoryTaskStore(), MemoryQueue()
    manager = TaskManager(store, queue, {})
    async with d_database.session() as session:
        service = UploadService(
            UploadRepository(session),
            storage,
            importer,
            registration,
            PermissionService(
                FakePermissions(
                    [
                        TableGrant(
                            name="ds_7.data",
                            all_columns={"id": "BIGINT"},
                            team_visible=True,
                            team_columns={"id"},
                        )
                    ]
                )
            ),
            AuditService(DatabaseAuditStore(session)),
            manager,
        )
        manager.handlers["upload.import"] = service.run_import
        yield service, manager, storage, importer, registration, session


def request(data=b"code,amount\n00123,2.5\n", datasource_id=None):
    return PreviewRequest(
        filename="orders.csv",
        content_base64=base64.b64encode(data).decode(),
        datasource_id=datasource_id,
    )


async def test_upload_preview_confirm_history_and_append(upload_setup):
    service, tasks, storage, importer, registration, session = upload_setup
    actor = CurrentUser(id=2, role=Role.LEADER, team_id=1)
    first = await service.preview(request(), actor=actor)
    assert first.columns[0].data_type == "text" and first.warnings
    assert first.rows[0][0] == "00123"
    job = await service.confirm(first.id, ConfirmRequest(config=first.config), actor=actor)
    await tasks.run(job.task_id)
    assert tasks.store.records[job.task_id].status == TaskStatus.SUCCESS
    completed = await service.get(job.id, actor=actor)
    assert completed.status == "success" and completed.row_count == 1
    assert registration.published[0][1].read_only
    assert registration.published[0][2].columns[0].comment.startswith("code")
    assert any(key.startswith("uploads/team_1/ds_7/v1/") for key in storage.values)
    assert not any(key.startswith("tmp/") for key in storage.values)
    second = await service.preview(request(b"code,amount\n00456,4.5\n", 7), actor=actor)
    job2 = await service.confirm(
        second.id, ConfirmRequest(config=second.config, mode="append"), actor=actor
    )
    await tasks.run(job2.task_id)
    assert len(importer.tables["data"][1]) == 2
    assert [job.version for job in await service.history(7, 0, 20, actor=actor)] == [2, 1]
    with pytest.raises(AppError):
        await service.confirm(first.id, ConfirmRequest(config=first.config), actor=actor)
    from sqlalchemy import select

    assert len(list(await session.scalars(select(AuditLog)))) == 6


async def test_registration_retry_does_not_repeat_append(upload_setup):
    service, tasks, _, importer, registration, _ = upload_setup
    actor = CurrentUser(id=2, role=Role.LEADER, team_id=1)
    preview = await service.preview(request(), actor=actor)
    registration.fail = True
    job = await service.confirm(preview.id, ConfirmRequest(config=preview.config), actor=actor)
    await tasks.run(job.task_id)
    assert (await service.get(job.id, actor=actor)).status == "registration_pending"
    registration.fail = False
    assert (await service.publish(job.id, actor=actor)).status == "success"
    assert importer.writes == 1


async def test_import_failure_preserves_previous_version(upload_setup):
    service, tasks, _, importer, _, _ = upload_setup
    actor = CurrentUser(id=2, role=Role.LEADER, team_id=1)
    preview = await service.preview(request(), actor=actor)
    job = await service.confirm(preview.id, ConfirmRequest(config=preview.config), actor=actor)
    await tasks.run(job.task_id)
    importer.fail = True
    preview2 = await service.preview(request(b"code,amount\n002,7.5\n", 7), actor=actor)
    failed = await service.confirm(preview2.id, ConfirmRequest(config=preview2.config), actor=actor)
    await tasks.run(failed.task_id)
    assert (await service.get(failed.id, actor=actor)).status == "failed"
    assert importer.tables["data"][1] == [("00123", 2.5)]


async def test_upload_roles_owner_and_append_difference(upload_setup):
    service, tasks, _, importer, _, _ = upload_setup
    leader = CurrentUser(id=2, role=Role.LEADER, team_id=1)
    with pytest.raises(AppError):
        await service.preview(request(), actor=CurrentUser(id=3, role=Role.MEMBER, team_id=1))
    preview = await service.preview(request(), actor=leader)
    with pytest.raises(AppError):
        await service.get(preview.id, actor=CurrentUser(id=4, role=Role.LEADER, team_id=2))
    job = await service.confirm(preview.id, ConfirmRequest(config=preview.config), actor=leader)
    await tasks.run(job.task_id)
    changed = await service.preview(request(b"new,amount\nx,1.0\n", 7), actor=leader)
    rejected = await service.confirm(
        changed.id, ConfirmRequest(config=changed.config, mode="append"), actor=leader
    )
    await tasks.run(rejected.task_id)
    assert tasks.store.records[rejected.task_id].status == TaskStatus.FAILED
    assert len(importer.tables["data"][1]) == 1


@pytest.mark.parametrize(
    "filename,data",
    [
        ("x.xlsm", b"data"),
        ("x.xlsx", b"code\n1"),
        ("x.xls", b"not-ole"),
        ("x.csv", b"PK\x03\x04binary"),
        ("../x.csv", b"a\n1"),
        ("x.csv", b"a\x00"),
    ],
)
def test_file_headers_rejected(filename, data):
    with pytest.raises(AppError):
        validate_file(filename, data)


def test_preview_cap_no_header_csv_encoding_and_excel_values():
    _, rows, _, _ = parse_rows("many.csv", b"id\n" + b"1\n" * 150, ParseConfig(), 100)
    assert len(rows) == 100
    headers, rows, _, _ = parse_rows(
        "budget.csv",
        "市场\t2025-02-01\t5.5\n".encode("gb18030"),
        ParseConfig(header_row=None, encoding="gb18030", delimiter="\t"),
    )
    assert len(headers) == 3 and rows[0][0] == "市场"
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "formula"])
    sheet.append(["00123", "=SUM(1,2)"])
    stream = BytesIO()
    book.save(stream)
    _, values, _, _ = parse_rows("x.xlsx", stream.getvalue(), ParseConfig())
    assert values == [["00123", None]]  # Formula cached value only; never evaluated.
    sheet.merge_cells("A3:B3")
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(AppError, match="请先整理"):
        parse_rows("x.xlsx", stream.getvalue(), ParseConfig())


def test_column_cleaning_duplicates_and_type_conversion():
    columns, warnings = configured_columns(
        ParseConfig(),
        ["价格(元)", "价格(元)", "code", "date"],
        [["1.5", "2", "00123", "2025-01-01"]],
    )
    assert len({c.name for c in columns}) == 4 and len(warnings) == 2
    with pytest.raises(AppError):
        convert_rows([["1.5"]], [ImportColumn(name="id", data_type="integer")])
    with pytest.raises(AppError):
        convert_rows([["NaN"]], [ImportColumn(name="amount", data_type="number")])


async def test_real_copy_gateway_batches_and_rolls_back_on_failure():
    class Connection:
        def __init__(self):
            self.calls, self.batches = [], []
            self.committed = self.rolled_back = self.closed = False

        @asynccontextmanager
        async def transaction(self):
            try:
                yield
            except Exception:
                self.rolled_back = True
                self.batches.clear()
                raise
            else:
                self.committed = True

        async def execute(self, sql, *args):
            self.calls.append(sql)

        async def fetch(self, sql, *args):
            return []

        async def fetchval(self, sql, *args):
            return False

        async def copy_records_to_table(self, name, **kwargs):
            self.batches.append(kwargs["records"][:])
            if len(self.batches) == 2:
                raise RuntimeError("private driver error")

        async def close(self):
            self.closed = True

    connection = Connection()

    async def connect(**kwargs):
        assert kwargs["user"] == "writer"
        return connection

    writer = Datasource(
        id=1,
        dialect=Dialect.POSTGRES,
        host="fake",
        port=5432,
        database="imports",
        username="writer",
        password=SecretStr("write-only"),
        read_only=False,
    )
    gateway = PostgresImportGateway(writer, SecretStr("read-only"), connect)
    plan = ImportPlan(
        datasource_id=7, table_name="data", columns=[ImportColumn(name="id")], batch_size=2
    )
    with pytest.raises(AppError):
        async with gateway.transaction(plan) as transaction:
            await transaction.write([("a",), ("b",), ("c",)])
    assert connection.rolled_back and not connection.committed and connection.closed
    assert connection.batches == []
    assert gateway.query_datasource(7).username == "ds_7_reader"
    assert gateway.query_datasource(7).read_only


def test_legacy_xls_preserves_string_codes():
    pytest.importorskip("xlrd", reason="Install upload/requirements.txt for legacy XLS")
    data = (Path(__file__).parent / "fixtures/upload_legacy.xls").read_bytes()
    headers, rows, sheets, _ = parse_rows("legacy.xls", data, ParseConfig())
    assert headers == ["code", "amount"] and sheets == ["Sales"]
    assert rows == [["00123", 2.5]]


@pytest.mark.parametrize("index", range(7))
def test_all_eval_upload_files_parse_with_manifest(index):
    import yaml

    manifest = yaml.safe_load((ROOT / "eval/files/manifest.yaml").read_text())
    job = manifest["jobs"][index]
    config = ParseConfig(
        sheet=job["sheet"] or 0,
        header_row=job["header_row"],
        encoding=job["encoding"],
        delimiter=job["delimiter"],
        table_name=job["table"],
        columns=[
            ImportColumn(
                name=c["name"],
                data_type=("number" if c["type"].startswith("numeric") else c["type"]),
            )
            for c in job["columns"]
        ],
    )
    headers, rows, _, _ = parse_rows(
        job["file"], (ROOT / "eval/files" / job["file"]).read_bytes(), config
    )
    columns, _ = configured_columns(config, headers, rows)
    converted = convert_rows(rows, columns)
    assert len(converted) == job["expected_rows"]
    if index == 0:
        assert converted[0][0] == "00123"


def test_api_import_gateway_has_only_reader_credentials():
    from app.core.importer import ImportReadGateway

    template = Datasource(
        id=1,
        dialect=Dialect.POSTGRES,
        host="pg-import",
        port=5432,
        database="imports",
        username="template",
        password=SecretStr("reader-only"),
    )
    gateway = ImportReadGateway(template)
    assert gateway.query_datasource(7).username == "ds_7_reader"
    assert gateway.query_datasource(7).password.get_secret_value() == "reader-only"
    with pytest.raises(AppError):
        gateway.transaction(
            ImportPlan(datasource_id=7, table_name="data", columns=[ImportColumn(name="id")])
        )
