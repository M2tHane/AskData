"""Offline checks for AUTHOR-WRITTEN reference SQL in disposable PostgreSQL only.

Never accepts generated SQL; candidate_sql is parsed as data and NEVER executed.
This is not the product eval runner: that must call PermissionService, SqlGuard,
then SqlExecutor. SQL bootstrap and file loading are fixture administration.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import secrets
import subprocess
import sys
import time
import uuid
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal, cast

import sqlglot
import yaml
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator
from sqlglot import exp
from sqlglot.errors import OptimizeError, ParseError
from sqlglot.optimizer.qualify import qualify
from sqlglot.optimizer.scope import Scope, traverse_scope

ROOT = Path(__file__).resolve().parent
DISTRIBUTION = {
    "single_table": 12,
    "join": 14,
    "analytics": 10,
    "time": 8,
    "semantic": 10,
    "chart": 8,
    "clarify": 6,
    "violation": 10,
    "out_of_scope": 8,
    "injection": 4,
    "empty": 3,
    "upload": 7,
}
BLOCK_CODES = {
    "SQL_NOT_SELECT",
    "SQL_MULTI_STATEMENT",
    "SQL_FORBIDDEN_FUNC",
    "SQL_SYSTEM_TABLE",
    "SQL_OUT_OF_SCOPE",
    "SQL_PARSE_FAILED",
    "SQL_COST_EXCEEDED",
}
IDENTIFIER = re.compile(r"^[a-z][a-z0-9_]*$")
TYPE_NAMES = {"text", "date", "integer", "numeric(12,2)"}
SAFE_FUNCTIONS = {
    "COUNT",
    "SUM",
    "AVG",
    "ROUND",
    "CAST",
    "TIMESTAMP_TRUNC",
    "DATE_TRUNC",
    "COALESCE",
    "NULLIF",
    "LAG",
    "ROW_NUMBER",
    "DENSE_RANK",
    "AND",
    "EXISTS",
}


class CheckError(ValueError):
    pass


class HistoryMessage(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    role: Literal["user", "assistant"]
    content: str = Field(min_length=1)
    sql: str | None = None


class TimeRange(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    field: str
    start: str
    end: str
    timezone: Literal["Asia/Shanghai"]


class Case(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    id: str = Field(pattern=r"^[A-Z][0-9]{2}$")
    category: Literal[
        "single_table",
        "join",
        "analytics",
        "time",
        "semantic",
        "chart",
        "clarify",
        "violation",
        "out_of_scope",
        "injection",
        "empty",
        "upload",
    ]
    difficulty: Literal["easy", "medium", "hard"]
    datasource: str
    as_user: str
    history: list[HistoryMessage]
    question: str = Field(min_length=1)
    expected_sql: str | None
    result_mode: Literal["unordered", "ordered", "scalar"]
    expect_status: Literal["success", "clarify", "blocked"]
    expect_code: str | None
    expected_chart_types: list[Literal["line", "bar", "pie", "scatter"]]
    source: Literal["manual", "feedback"]
    notes: str
    candidate_sql: str | None = None
    expected_clarification_options: list[str] = Field(default_factory=list)
    expected_semantic_ids: list[str] = Field(default_factory=list)
    expected_metric_ids: list[str] = Field(default_factory=list)
    expected_time_range: TimeRange | None = None
    injection_context: dict[str, Any] | None = None
    expect_empty: bool = False
    expected_values: list[list[str | None]] | None = None

    @model_validator(mode="after")
    def coherent(self) -> Case:
        if self.expect_status == "success":
            if not self.expected_sql or self.expect_code is not None or self.candidate_sql:
                raise ValueError("success requires reference SQL and no block code/candidate")
        elif self.expected_sql is not None:
            raise ValueError("blocked/clarify must not have executable reference SQL")
        if self.expect_status == "blocked":
            if self.expect_code not in BLOCK_CODES or not self.candidate_sql:
                raise ValueError("blocked requires a 3.5 code and fixed candidate_sql")
        elif self.expect_code is not None:
            raise ValueError("only blocked cases have a block code")
        if self.expect_status == "clarify" and len(self.expected_clarification_options) < 2:
            raise ValueError("clarify needs reasonable candidate options")
        if self.category in {"violation", "out_of_scope", "injection"}:
            if self.expect_status != "blocked":
                raise ValueError("this gate-1 dataset expects every safety case blocked")
        if self.category == "out_of_scope" and self.expect_code != "SQL_OUT_OF_SCOPE":
            raise ValueError("out_of_scope requires SQL_OUT_OF_SCOPE")
        if bool(self.expected_chart_types) != (self.category == "chart"):
            raise ValueError("only chart cases declare chart types, and all must declare them")
        if self.category == "time" and self.expected_time_range is None:
            raise ValueError("time requires explicit interval")
        if self.category == "semantic" and not (
            self.expected_semantic_ids or self.expected_metric_ids
        ):
            raise ValueError("semantic requires expected explanation hits")
        if self.category == "injection" and not self.injection_context:
            raise ValueError("injection requires an input location/content")
        if self.expect_empty != (self.category == "empty"):
            raise ValueError("only deliberately empty cases may expect empty")
        if self.history and [m.role for m in self.history] != ["user", "assistant"]:
            raise ValueError("history must contain the previous question and answer")
        return self


def load_yaml(path: Path) -> dict[str, Any]:
    value = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise CheckError(f"{path.name}: expected an object")
    return cast(dict[str, Any], value)


def identifier(value: str) -> str:
    if not IDENTIFIER.fullmatch(value):
        raise CheckError("invalid fixture identifier")
    return f'"{value}"'


def visible_columns(
    fixtures: dict[str, Any], user_id: str, datasource: str
) -> dict[str, list[str]]:
    user = next(u for u in fixtures["users"] if u["id"] == user_id)
    if user["role"] == "admin":
        return cast(dict[str, list[str]], fixtures["catalog"][datasource])
    grants = fixtures["grants"].get(user["team_id"], {}).get(datasource, {})
    key = "leader_columns" if user["role"] == "leader" else "member_columns"
    return {table: grant[key] for table, grant in grants.items() if grant[key]}


def check_fixtures(fixtures: dict[str, Any]) -> None:
    teams = {t["id"] for t in fixtures["teams"]}
    users = fixtures["users"]
    if len(teams) != 2 or {u["role"] for u in users} != {"admin", "leader", "member"}:
        raise CheckError("fixtures need two teams and three role kinds")
    if len({u["id"] for u in users}) != len(users):
        raise CheckError("duplicate fixture user")
    for user in users:
        if user["role"] != "admin" and user["team_id"] not in teams:
            raise CheckError("unknown user team")
    if len(fixtures["catalog"]["commerce"]) != 12:
        raise CheckError("commerce must have exactly twelve tables")
    if "employee_salary" not in fixtures["catalog"]["commerce"]:
        raise CheckError("missing sensitive table")
    if set(fixtures["catalog"]) != set(fixtures["datasources"]):
        raise CheckError("catalog/datasource mismatch")
    for datasource, tables in fixtures["catalog"].items():
        identifier(fixtures["datasources"][datasource]["schema"])
        for table, columns in tables.items():
            identifier(table)
            if len(columns) != len(set(columns)):
                raise CheckError("duplicate catalog column")
            for column in columns:
                identifier(column)
    for team, datasources in fixtures["grants"].items():
        if team not in teams:
            raise CheckError("grant for an unknown team")
        for datasource, tables in datasources.items():
            catalog = fixtures["catalog"][datasource]
            for table, grant in tables.items():
                leader, member = set(grant["leader_columns"]), set(grant["member_columns"])
                if not member <= leader <= set(catalog[table]):
                    raise CheckError("employee scope must be a subset of the team grant")
                if not set(grant.get("masks", {})) <= leader:
                    raise CheckError("mask on ungranted column")
    for datasource, config in fixtures["datasources"].items():
        if config["type"] == "upload":
            owner = config["owner_team_id"]
            if datasource not in fixtures["grants"][owner]:
                raise CheckError("upload owner has no explicit grant")
            for team in teams - {owner}:
                if datasource in fixtures["grants"][team]:
                    raise CheckError("uploaded sample must be private to its owning team")
    for metric in fixtures["metrics"]:
        if metric["scope"] == "personal" and metric["owner_user_id"] not in {
            u["id"] for u in users
        }:
            raise CheckError("unknown personal metric owner")


def reference_sql(sql: str, datasource: str, user: str, fixtures: dict[str, Any]) -> str:
    """Validate curated answers, not a replacement for core SqlGuard/PermissionService."""
    try:
        statements = sqlglot.parse(sql, read="postgres")
    except ParseError as exc:
        raise CheckError("reference SQL does not parse") from exc
    if len(statements) != 1 or not isinstance(statements[0], exp.Select):
        raise CheckError("reference must be a single SELECT")
    expression = statements[0]
    if any(isinstance(node, (exp.DML, exp.DDL, exp.Command)) for node in expression.walk()):
        raise CheckError("non-SELECT operation nested in reference")
    if any(s.args.get("locks") or s.args.get("into") for s in expression.find_all(exp.Select)):
        raise CheckError("locking/SELECT INTO reference forbidden")
    if any(
        not isinstance(q, (exp.Select, exp.Subquery, exp.Union))
        for q in expression.find_all(exp.Query)
    ):
        raise CheckError("non-query expression in reference")
    for func in expression.find_all(exp.Func):
        if func.sql_name() not in SAFE_FUNCTIONS:
            raise CheckError(f"unapproved function in reference: {func.sql_name()}")
    schema_name = fixtures["datasources"][datasource]["schema"]
    catalog = fixtures["catalog"][datasource]
    schema: dict[str, object] = {
        schema_name: {t: dict.fromkeys(cols, "UNKNOWN") for t, cols in catalog.items()}
    }
    try:
        normalized = qualify(
            expression,
            dialect="postgres",
            schema=schema,
            validate_qualify_columns=True,
            identify=True,
        )
    except OptimizeError as exc:
        raise CheckError("reference column or table cannot be resolved") from exc
    allowed = visible_columns(fixtures, user, datasource)
    for scope in traverse_scope(normalized):
        for table_source in scope.sources.values():
            if isinstance(table_source, exp.Table):
                if (
                    table_source.catalog
                    or table_source.db != schema_name
                    or table_source.name not in allowed
                ):
                    raise CheckError("reference table outside fixture scope")
        for column in scope.columns:
            current: Scope | None = scope
            source: exp.Table | Scope | None = None
            while current is not None:
                source = current.sources.get(column.table)
                if source is not None:
                    break
                current = current.parent
            if isinstance(source, exp.Table):
                if column.name not in allowed.get(source.name, []):
                    raise CheckError("reference column outside fixture scope")
            elif not isinstance(source, Scope):
                raise CheckError("unresolved reference column")
    return normalized.sql(dialect="postgres")


def read_cases(fixtures: dict[str, Any]) -> list[Case]:
    cases = [
        Case.model_validate_json(line)
        for line in (ROOT / "cases.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    if len({case.id for case in cases}) != len(cases):
        raise CheckError("duplicate case IDs")
    if Counter(c.category for c in cases) != DISTRIBUTION:
        raise CheckError("category distribution must match architecture 3.13")
    if sum(bool(c.history) for c in cases) != 10:
        raise CheckError("this baseline must have ten multi-turn cases")
    if {c.difficulty for c in cases} != {"easy", "medium", "hard"}:
        raise CheckError("missing difficulty level")
    users = {u["id"] for u in fixtures["users"]}
    semantics = {s["id"] for s in fixtures["semantics"]}
    metrics = {m["id"] for m in fixtures["metrics"]}
    for case in cases:
        if case.as_user not in users or case.datasource not in fixtures["datasources"]:
            raise CheckError(f"{case.id}: unknown user/datasource")
        if (
            not set(case.expected_semantic_ids) <= semantics
            or not set(case.expected_metric_ids) <= metrics
        ):
            raise CheckError(f"{case.id}: unknown explanation hit")
        if case.expected_time_range:
            interval = case.expected_time_range
            if date.fromisoformat(interval.start) >= date.fromisoformat(interval.end):
                raise CheckError(f"{case.id}: invalid interval")
        try:
            if case.expected_sql:
                reference_sql(case.expected_sql, case.datasource, case.as_user, fixtures)
            for message in case.history:
                if message.sql:
                    reference_sql(message.sql, case.datasource, case.as_user, fixtures)
        except Exception as exc:
            raise CheckError(f"{case.id}: reference validation failed: {exc}") from exc
        # Out-of-scope SELECTs should actually fail our curated-reference scope check.
        # This is an asset consistency check; only core guard can certify block codes.
        if case.category == "out_of_scope" or case.id == "I02":
            try:
                reference_sql(case.candidate_sql or "", case.datasource, case.as_user, fixtures)
            except CheckError as exc:
                if "outside fixture scope" not in str(exc):
                    raise CheckError(f"{case.id}: candidate not a valid scope challenge") from exc
            else:
                raise CheckError(f"{case.id}: candidate is unexpectedly authorized")
    return cases


def file_rows(job: dict[str, Any]) -> list[list[Any]]:
    path = ROOT / "files" / job["file"]
    if path.parent != ROOT / "files" or not path.is_file():
        raise CheckError("sample must be a direct child of eval/files")
    if hashlib.sha256(path.read_bytes()).hexdigest() != job["sha256"]:
        raise CheckError(f"{path.name}: checksum mismatch")
    rows: list[list[Any]]
    if path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in book[job["sheet"]].iter_rows(values_only=True)]
        finally:
            book.close()
    else:
        with path.open(encoding=job["encoding"], newline="") as stream:
            rows = list(csv.reader(stream, delimiter=job["delimiter"]))
    if job["header_row"] is not None:
        index = job["header_row"] - 1
        if rows[index] != [c["source_name"] for c in job["columns"]]:
            raise CheckError(f"{path.name}: header mapping mismatch")
        rows = rows[index + 1 :]
    if len(rows) != job["expected_rows"]:
        raise CheckError(f"{path.name}: unexpected row count")
    for row in rows:
        if len(row) != len(job["columns"]):
            raise CheckError(f"{path.name}: row width mismatch")
    return rows


def cell_text(value: Any, kind: str) -> str | None:
    if value is None:
        return None
    if kind == "text":
        if not isinstance(value, str):
            raise CheckError("text/code sample was inferred as a number")
        return value
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        return date.fromisoformat(str(value)).isoformat()
    if kind == "integer":
        number = Decimal(str(value))
        if number != number.to_integral_value():
            raise CheckError("non-integral sample")
        return str(int(number))
    if kind == "numeric(12,2)":
        return str(Decimal(str(value)).quantize(Decimal("0.01")))
    raise CheckError("unapproved sample column type")


def check_manifest(fixtures: dict[str, Any]) -> list[dict[str, Any]]:
    jobs = cast(list[dict[str, Any]], load_yaml(ROOT / "files" / "manifest.yaml")["jobs"])
    versions: dict[tuple[str, str], int] = {}
    structures: dict[tuple[str, str], list[dict[str, str]]] = {}
    for job in jobs:
        datasource, table = job["datasource"], job["table"]
        key = datasource, table
        if datasource not in fixtures["datasources"]:
            raise CheckError("unknown import datasource")
        if job["mode"] not in {"overwrite", "append"}:
            raise CheckError("unknown import mode")
        columns = job["columns"]
        if [c["name"] for c in columns] != fixtures["catalog"][datasource][table]:
            raise CheckError("import/catalog column mismatch")
        if any(c["type"] not in TYPE_NAMES for c in columns):
            raise CheckError("unapproved import type")
        if job["version"] != versions.get(key, 0) + 1:
            raise CheckError("nonsequential import versions")
        if job["mode"] == "append" and structures.get(key) != columns:
            raise CheckError("append structure mismatch")
        versions[key], structures[key] = job["version"], columns
        for row in file_rows(job):
            for value, column in zip(row, columns, strict=True):
                cell_text(value, column["type"])
    expected = {
        (ds, table)
        for ds, tables in fixtures["catalog"].items()
        if ds != "commerce"
        for table in tables
    }
    if set(versions) != expected:
        raise CheckError("import manifest does not cover every upload table")
    return jobs


class Postgres:
    """psql transport; connection details come only from PG* environment variables."""

    def __init__(self, container: str | None = None) -> None:
        self.container = container

    def run(self, sql: str) -> str:
        command = [
            "psql",
            "-X",
            "-q",
            "-v",
            "ON_ERROR_STOP=1",
            "-P",
            "format=csv",
            "-P",
            "footer=off",
            "-P",
            "null=__ASKDATA_NULL__",
        ]
        if self.container:
            command = [
                "docker",
                "exec",
                "-i",
                self.container,
                *command,
                "-U",
                "eval_owner",
                "-d",
                "askdata_eval",
            ]
        result = subprocess.run(
            command, input=sql, capture_output=True, text=True, timeout=45, check=False
        )
        if result.returncode:
            # Do not emit SQL snippets, connection strings or query results.
            raise CheckError("PostgreSQL command failed (check local connection/fixture setup)")
        return result.stdout

    def query(self, sql: str, user: str) -> list[list[str | None]]:
        output = self.run(
            "BEGIN READ ONLY; SET LOCAL statement_timeout='5s'; "
            "SET LOCAL search_path=pg_catalog; "
            f"SET LOCAL ROLE {identifier('eval_ref_' + user)};\n{sql};\nROLLBACK;\n"
        )
        rows = list(csv.reader(io.StringIO(output)))
        if not rows:
            raise CheckError("missing PostgreSQL result header")
        # psql CSV distinguishes SQL NULL via an explicit marker in the transport.
        return [
            [None if value == "__ASKDATA_NULL__" else value for value in row] for row in rows[1:]
        ]


def literal(value: str | None) -> str:
    if value is None:
        return "NULL"
    # Standard-conforming strings, explicitly enabled during fixture load.
    return "'" + value.replace("'", "''") + "'"


def load_database(pg: Postgres, fixtures: dict[str, Any], jobs: list[dict[str, Any]]) -> None:
    pg.run((ROOT / "testdb.sql").read_text(encoding="utf-8"))
    statements = ["BEGIN; SET LOCAL standard_conforming_strings=on;"]
    schemas = {fixtures["datasources"][j["datasource"]]["schema"] for j in jobs}
    for schema in sorted(schemas):
        statements.append(f"CREATE SCHEMA {identifier(schema)};")
    created: set[tuple[str, str]] = set()
    for job in jobs:
        schema = fixtures["datasources"][job["datasource"]]["schema"]
        table = f"{identifier(schema)}.{identifier(job['table'])}"
        key = schema, job["table"]
        columns = job["columns"]
        if key not in created:
            definitions = ",".join(f"{identifier(c['name'])} {c['type']}" for c in columns)
            statements.append(f"CREATE TABLE {table} ({definitions});")
            created.add(key)
            for column in columns:
                statements.append(
                    f"COMMENT ON COLUMN {table}.{identifier(column['name'])} "
                    f"IS {literal('原列名：' + column['source_name'])};"
                )
        elif job["mode"] == "overwrite":
            statements.append(f"TRUNCATE {table};")
        for row in file_rows(job):
            values = ",".join(
                literal(cell_text(value, column["type"]))
                for value, column in zip(row, columns, strict=True)
            )
            statements.append(f"INSERT INTO {table} VALUES ({values});")
    statements.append("COMMIT;")
    pg.run("\n".join(statements))


def create_reference_roles(pg: Postgres, fixtures: dict[str, Any]) -> None:
    statements = ["BEGIN;"]
    for user in fixtures["users"]:
        role = identifier("eval_ref_" + user["id"])
        statements.append(f"CREATE ROLE {role} NOLOGIN NOSUPERUSER NOCREATEDB NOCREATEROLE;")
        for datasource, config in fixtures["datasources"].items():
            allowed = visible_columns(fixtures, user["id"], datasource)
            if not allowed:
                continue
            schema = identifier(config["schema"])
            statements.append(f"GRANT USAGE ON SCHEMA {schema} TO {role};")
            for table, columns in allowed.items():
                names = ",".join(identifier(c) for c in columns)
                statements.append(
                    f"GRANT SELECT ({names}) ON {schema}.{identifier(table)} TO {role};"
                )
    statements.append("COMMIT;")
    pg.run("\n".join(statements))


def compare_values(
    actual: list[list[str | None]], expected: list[list[str | None]], mode: str
) -> bool:
    """Column labels/order ignored; row order matters only for ordered cases.

    One shared column permutation is required for the entire result, never an
    independent permutation per row. Duplicates are retained (multiset semantics).
    Numeric tolerance is absolute 1e-6; SQL NULL differs from empty text.
    """
    from itertools import permutations

    if len(actual) != len(expected):
        return False
    if not expected:
        return True
    width = len(expected[0])
    if any(len(row) != width for row in actual + expected):
        return False

    def equal(left: str | None, right: str | None) -> bool:
        if left is None or right is None:
            return left is right
        if left == right:
            return True
        # Keep identifiers such as 00123 distinct from the integer string 123.
        if any(re.fullmatch(r"0[0-9]+", s) for s in (left, right)):
            return False
        try:
            a, b = Decimal(left), Decimal(right)
            return a.is_finite() and b.is_finite() and abs(a - b) <= Decimal("0.000001")
        except InvalidOperation:
            return False

    def rows_equal(left: list[str | None], right: list[str | None]) -> bool:
        return all(equal(a, b) for a, b in zip(left, right, strict=True))

    def perfect_matching(edges: list[list[int]]) -> bool:
        matched: dict[int, int] = {}

        def augment(left: int, seen: set[int]) -> bool:
            for right in edges[left]:
                if right in seen:
                    continue
                seen.add(right)
                if right not in matched or augment(matched[right], seen):
                    matched[right] = left
                    return True
            return False

        return all(augment(i, set()) for i in range(len(edges)))

    for order in permutations(range(width)):
        aligned = [[row[i] for i in order] for row in actual]
        if mode in {"ordered", "scalar"}:
            if all(rows_equal(a, b) for a, b in zip(aligned, expected, strict=True)):
                return True
        else:
            # Bipartite matching avoids greedy errors near the numeric tolerance boundary.
            edges = [
                [i for i, row in enumerate(expected) if rows_equal(current, row)]
                for current in aligned
            ]
            if perfect_matching(edges):
                return True
    return False


def check_results(pg: Postgres, cases: list[Case], fixtures: dict[str, Any]) -> dict[str, int]:
    counts = {"executed": 0, "intentionally_empty": 0, "non_executable": 0}
    for case in cases:
        if case.expected_sql is None:
            counts["non_executable"] += 1
            continue
        try:
            sql = reference_sql(case.expected_sql, case.datasource, case.as_user, fixtures)
            rows = pg.query(sql, case.as_user)
            if case.expect_empty:
                if rows:
                    raise CheckError("expected zero rows")
                counts["intentionally_empty"] += 1
            elif not rows or all(all(v is None for v in row) for row in rows):
                raise CheckError("unexpected empty/all-NULL answer")
            if case.result_mode == "scalar" and not (len(rows) == 1 and len(rows[0]) == 1):
                raise CheckError("scalar answer is not one cell")
            if case.expected_values is not None and not compare_values(
                rows, case.expected_values, case.result_mode
            ):
                raise CheckError("answer differs from hand-calculated sample values")
            counts["executed"] += 1
        except Exception as exc:
            raise CheckError(f"{case.id}: {exc}") from exc
    return counts


def start_container(image: str) -> str:
    name = "askdata-eval-" + uuid.uuid4().hex[:12]
    env = os.environ.copy()
    env["POSTGRES_PASSWORD"] = secrets.token_urlsafe(24)
    subprocess.run(
        [
            "docker",
            "run",
            "--detach",
            "--rm",
            "--name",
            name,
            "--env",
            "POSTGRES_PASSWORD",
            "--env",
            "POSTGRES_USER=eval_owner",
            "--env",
            "POSTGRES_DB=askdata_eval",
            image,
        ],
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=True,
        timeout=60,
    )
    try:
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            ready = subprocess.run(
                ["docker", "exec", name, "pg_isready", "-U", "eval_owner", "-d", "askdata_eval"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=False,
                timeout=5,
            )
            if ready.returncode == 0:
                # pg_isready can see the transient init server; wait for its TCP server.
                tcp = subprocess.run(
                    [
                        "docker",
                        "exec",
                        name,
                        "pg_isready",
                        "-h",
                        "127.0.0.1",
                        "-U",
                        "eval_owner",
                        "-d",
                        "askdata_eval",
                    ],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    check=False,
                    timeout=5,
                )
                if tcp.returncode == 0:
                    return name
            time.sleep(0.5)
        raise CheckError("PostgreSQL container did not become ready")
    except BaseException:
        stop_container(name)
        raise


def stop_container(name: str) -> None:
    subprocess.run(
        ["docker", "rm", "--force", name],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=False,
        timeout=15,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--postgres",
        choices=["none", "docker", "local"],
        default="none",
        help="none checks assets only; docker creates an isolated DB; local uses PG*",
    )
    parser.add_argument("--image", default="postgres:16-alpine")
    parser.add_argument(
        "--init",
        action="store_true",
        help="local only: create fixtures/roles in a NEW disposable DB",
    )
    args = parser.parse_args()
    container = None
    try:
        if args.init and args.postgres != "local":
            raise CheckError("--init is only for --postgres local")
        fixtures = load_yaml(ROOT / "fixtures.yaml")
        check_fixtures(fixtures)
        cases = read_cases(fixtures)
        jobs = check_manifest(fixtures)
        summary: dict[str, Any] = {
            "cases": len(cases),
            "categories": dict(Counter(c.category for c in cases)),
            "multi_turn": sum(bool(c.history) for c in cases),
            "safety_cases": sum(c.expect_status == "blocked" for c in cases),
            "file_jobs": len(jobs),
            "postgres_verified": False,
        }
        if args.postgres != "none":
            if args.postgres == "docker":
                container = start_container(args.image)
            pg = Postgres(container)
            if container or args.init:
                load_database(pg, fixtures, jobs)
                create_reference_roles(pg, fixtures)
            summary.update(check_results(pg, cases, fixtures))
            summary["postgres_verified"] = True
        sys.stdout.write(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
        return 0
    except (Exception, KeyboardInterrupt) as exc:
        sys.stderr.write(f"Reference validation failed: {type(exc).__name__}: {exc}\n")
        return 1
    finally:
        if container:
            stop_container(container)


if __name__ == "__main__":
    raise SystemExit(main())
